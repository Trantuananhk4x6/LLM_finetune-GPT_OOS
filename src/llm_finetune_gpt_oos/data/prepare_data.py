import csv
import hashlib
import json
import math
import random
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from ..config import load_config
from .dataset import write_jsonl


SYSTEM_PROMPT = "Bạn là Sales Operations AI và tư vấn viên bán hàng chuyên nghiệp. Chỉ suy luận từ dữ liệu được cung cấp, nêu rõ khi dữ liệu không đủ, không bịa giá hoặc chính sách và không tiết lộ thông tin nhận dạng cá nhân. Trả lời ngắn gọn, lịch sự, có cấu trúc và ưu tiên hành động có thể kiểm chứng."
PII_EXACT_COLUMNS = {"name", "fullname", "firstname", "lastname", "customername", "customerid", "email", "phone", "phonenumber", "streetaddress", "address", "zipcode", "postalcode", "latitude", "longitude", "password", "birthdate", "dateofbirth", "ssn"}
PII_CONTAINS_TOKENS = {"email", "phone", "address", "street", "zipcode", "postalcode", "latitude", "longitude", "password", "birthdate", "ssn"}
CANONICAL_ALIASES = {
    "order_id": {"orderid", "order_id", "invoice", "invoiceno", "transactionid", "transaction_id", "bookingid", "booking_id"},
    "date": {"orderdate", "order_date", "transactiondate", "transaction_date", "date", "createdat", "created_at"},
    "product": {"product", "productname", "product_name", "item", "description", "productcategory", "product_category"},
    "category": {"category", "productcategory", "product_category", "mastercategory", "master_category"},
    "quantity": {"quantity", "qty", "quantitysold", "quantity_sold", "items"},
    "unit_price": {"price", "unitprice", "unit_price"},
    "amount": {"sales", "totalamount", "total_amount", "totalprice", "total_price", "revenue", "profit"},
    "discount": {"discount", "discountapplied", "discount_applied", "discountpercent", "discount_percent", "promoamount", "promo_amount"},
    "channel": {"channel", "storelocation", "store_location", "paymentmethod", "payment_method", "source"},
    "status": {"status", "paymentstatus", "payment_status", "resolutionstatus", "resolution_status", "shippingstatus", "shipping_status"},
    "issue": {"issuecategory", "issue_category", "notes", "reviewtext", "review_text", "interactiontype", "interaction_type"},
    "priority": {"priority", "customersatisfactionscore", "customer_satisfaction_score", "rating"},
}


def normalize_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def is_pii_column(column: str) -> bool:
    normalized = normalize_name(column)
    return normalized in PII_EXACT_COLUMNS or any(token in normalized for token in PII_CONTAINS_TOKENS)


def redact_text(value: str) -> str:
    value = re.sub(r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b", "[REDACTED_EMAIL]", value)
    return re.sub(r"(?<!\w)(?:\+?\d[\d .()-]{7,}\d)(?!\w)", "[REDACTED_PHONE]", value)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        cleaned = redact_text(" ".join(value.split()))
        return cleaned[:500] if cleaned else None
    if isinstance(value, float):
        return round(value, 4)
    return value


def to_number(value: Any) -> float | None:
    try:
        return float(str(value).replace(",", "").replace("%", ""))
    except (TypeError, ValueError):
        return None


def find_column(columns: Iterable[str], canonical: str) -> str | None:
    aliases = CANONICAL_ALIASES[canonical]
    return next((column for column in columns if normalize_name(column) in aliases), None)


def safe_record(row: dict[str, Any], redact_pii: bool) -> dict[str, Any]:
    return {
        key: normalized
        for key, value in row.items()
        if not (redact_pii and is_pii_column(key))
        if (normalized := normalize_value(value)) is not None
    }


def reservoir_sample(rows: Iterable[dict[str, Any]], limit: int, seed: int) -> tuple[list[dict[str, Any]], bool]:
    randomizer = random.Random(seed)
    sample: list[dict[str, Any]] = []
    count = 0
    for row in rows:
        count += 1
        if len(sample) < limit:
            sample.append(row)
            continue
        replacement = randomizer.randrange(count)
        if replacement < limit:
            sample[replacement] = row
    return sample, count > limit


def read_csv(path: Path, limit: int, seed: int) -> tuple[list[dict[str, Any]], bool]:
    with path.open("r", encoding="utf-8-sig", newline="", errors="replace") as handle:
        return reservoir_sample(csv.DictReader(handle), limit, seed)


def read_xlsx(path: Path, limit: int, seed: int) -> list[tuple[str, list[dict[str, Any]], bool]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    tables = []
    for worksheet in workbook.worksheets:
        iterator = worksheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            continue
        columns = [str(value).strip() if value is not None else f"column_{index}" for index, value in enumerate(header)]
        rows = (dict(zip(columns, values)) for values in iterator)
        sample, truncated = reservoir_sample(rows, limit, seed + len(tables))
        tables.append((f"{path.stem}_{worksheet.title}", sample, truncated))
    workbook.close()
    return tables


def extract_archives(input_dir: Path, extraction_dir: Path) -> list[Path]:
    extraction_dir.mkdir(parents=True, exist_ok=True)
    targets = []
    for archive in input_dir.glob("*.zip"):
        archive_digest = sha256_file(archive)[:12]
        target = extraction_dir / f"{archive.stem}-{archive_digest}"
        target.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(archive) as package:
            for member in package.infolist():
                destination = (target / member.filename).resolve()
                if target.resolve() not in destination.parents and destination != target.resolve():
                    raise ValueError(f"Archive contains an unsafe path: {member.filename}")
                package.extract(member, target)
        targets.append(target)
    return targets


def records_from_sources(input_dir: Path, extraction_dir: Path, max_rows_per_table: int, redact_pii: bool, seed: int) -> tuple[list[tuple[str, dict[str, Any]]], list[dict[str, Any]]]:
    extracted_dirs = extract_archives(input_dir, extraction_dir)
    csv_paths = sorted({*input_dir.rglob("*.csv"), *(path for directory in extracted_dirs for path in directory.rglob("*.csv"))})
    xlsx_paths = sorted({*input_dir.rglob("*.xlsx"), *(path for directory in extracted_dirs for path in directory.rglob("*.xlsx"))})
    records: list[tuple[str, dict[str, Any]]] = []
    lineage: list[dict[str, Any]] = []
    for path in csv_paths:
        sample, truncated = read_csv(path, max_rows_per_table, seed)
        digest = sha256_file(path)
        lineage.append({"file": str(path), "sha256": digest, "rows_used": len(sample), "truncated": truncated, "columns": list(sample[0]) if sample else []})
        records.extend((path.stem, safe_record(row, redact_pii)) for row in sample)
    for path in xlsx_paths:
        digest = sha256_file(path)
        for table_name, sample, truncated in read_xlsx(path, max_rows_per_table, seed):
            lineage.append({"file": str(path), "table": table_name, "sha256": digest, "rows_used": len(sample), "truncated": truncated, "columns": list(sample[0]) if sample else []})
            records.extend((table_name, safe_record(row, redact_pii)) for row in sample)
    return records, lineage


def transaction_answer(record: dict[str, Any]) -> str:
    columns = record.keys()
    quantity = to_number(record.get(find_column(columns, "quantity"), 1)) or 1
    price_key = find_column(columns, "unit_price")
    price = to_number(record.get(price_key)) if price_key else None
    amount_key = find_column(columns, "amount")
    amount = to_number(record.get(amount_key)) if amount_key else None
    discount_key = find_column(columns, "discount")
    discount = to_number(record.get(discount_key)) if discount_key else None
    parts = ["Kết luận", "- Đây là một bản ghi giao dịch cần được đối chiếu với KPI tổng hợp, không phải bằng chứng cho một xu hướng toàn hệ thống."]
    if price is not None:
        gross = quantity * price
        net = gross * (1 - min(max(discount or 0, 0), 100) / 100)
        parts.append(f"- Giá trị ước tính: số lượng {quantity:g} × đơn giá {price:,.2f} = {gross:,.2f}; sau chiết khấu ước tính {net:,.2f}.")
    elif amount is not None:
        parts.append(f"- Giá trị được ghi nhận trong bản ghi: {amount:,.2f}.")
    product_key = find_column(columns, "product")
    channel_key = find_column(columns, "channel")
    if product_key and record.get(product_key):
        parts.append(f"- Sản phẩm hoặc nhóm hàng: {record[product_key]}.")
    if channel_key and record.get(channel_key):
        parts.append(f"- Kênh hoặc điểm chạm: {record[channel_key]}.")
    parts.append("- Hành động tiếp theo: tổng hợp cùng kỳ theo sản phẩm, kênh và thời gian trước khi thay đổi giá hoặc ngân sách.")
    return "\n".join(parts)


def support_answer(record: dict[str, Any]) -> str:
    issue_key = find_column(record.keys(), "issue")
    priority_key = find_column(record.keys(), "priority")
    status_key = find_column(record.keys(), "status")
    parts = ["Đề xuất xử lý", "- Xác nhận định danh đơn hàng qua kênh bảo mật trước khi tra cứu chi tiết."]
    if issue_key and record.get(issue_key):
        parts.append(f"- Chủ đề ghi nhận: {record[issue_key]}.")
    if priority_key and record.get(priority_key):
        parts.append(f"- Mức ưu tiên hoặc chỉ số hài lòng: {record[priority_key]}.")
    if status_key and record.get(status_key):
        parts.append(f"- Trạng thái hiện có: {record[status_key]}.")
    parts.append("- Hành động tiếp theo: ghi nhận owner, SLA và mốc cập nhật cho khách hàng; chỉ đóng khi có xác nhận kết quả.")
    return "\n".join(parts)


def build_examples(source: str, record: dict[str, Any], examples_per_row: int) -> list[dict[str, Any]]:
    customer_prompt = record.get("customer_prompt")
    agent_response = record.get("agent_response")
    if customer_prompt and agent_response:
        identifier = hashlib.sha256(f"{source}|{customer_prompt}|{agent_response}".encode("utf-8")).hexdigest()
        return [{"id": identifier, "source_table": source, "messages": [{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": customer_prompt}, {"role": "assistant", "content": agent_response}]}]
    compact = json.dumps(record, ensure_ascii=False, separators=(",", ":"))
    issue_key = find_column(record.keys(), "issue")
    answer = support_answer(record) if issue_key else transaction_answer(record)
    requests = ["Phân tích bản ghi bán hàng dưới đây và nêu hành động tiếp theo có thể kiểm chứng.", "Hãy tóm tắt ý nghĩa vận hành của dữ liệu dưới đây. Không suy đoán ngoài dữ liệu."]
    examples = []
    for request in requests[:examples_per_row]:
        identifier = hashlib.sha256(f"{source}|{compact}|{request}".encode("utf-8")).hexdigest()
        examples.append({"id": identifier, "source_table": source, "messages": [{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": f"{request}\nDữ liệu: {compact}"}, {"role": "assistant", "content": answer}]})
    return examples


def validate_examples(rows: list[dict[str, Any]], minimum_characters: int) -> list[dict[str, Any]]:
    valid = []
    for row in rows:
        messages = row.get("messages", [])
        if len(messages) != 3 or [message.get("role") for message in messages] != ["system", "user", "assistant"]:
            continue
        if len(messages[-1].get("content", "")) < minimum_characters:
            continue
        valid.append(row)
    return valid


def prepare_data(config_path: str = "configs/dataset.yaml") -> dict[str, Any]:
    config = load_config(config_path)
    input_dir = Path(config["input_dir"])
    extraction_dir = Path(config["extraction_dir"])
    output_dir = Path(config["output_dir"])
    if not input_dir.exists():
        raise FileNotFoundError(f"Input directory does not exist: {input_dir}")
    records, lineage = records_from_sources(input_dir, extraction_dir, config["sampling"]["max_rows_per_table"], config["quality"]["redact_pii"], config["split"]["seed"])
    generated = [example for source, record in records for example in build_examples(source, record, config["sampling"]["examples_per_row"])]
    examples = validate_examples(generated, config["quality"]["min_assistant_characters"])
    if not examples:
        raise ValueError("No valid examples were generated from the supplied source files")
    random.Random(config["split"]["seed"]).shuffle(examples)
    boundary = int(len(examples) * config["split"]["train_ratio"])
    train_rows, validation_rows = examples[:boundary], examples[boundary:]
    output_dir.mkdir(parents=True, exist_ok=True)
    write_jsonl(output_dir / "train.jsonl", train_rows)
    write_jsonl(output_dir / "validation.jsonl", validation_rows)
    manifest = {"source": config["source"], "created_at": datetime.now().astimezone().isoformat(), "examples": len(examples), "train_examples": len(train_rows), "validation_examples": len(validation_rows), "source_tables": Counter(row["source_table"] for row in examples), "lineage": lineage}
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest
